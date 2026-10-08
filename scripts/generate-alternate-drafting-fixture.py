#!/usr/bin/env python3
"""Reproduce a separate synthetic Drafting correspondence and its sealed oracle.

Run with the Codex bundled Python (python-docx). This script never calls ConSense
or an LLM. Only 01_upload_documents contains model inputs. Every expected answer
is backed by an exact native paragraph and DOCX SHA-256 in the sealed reference.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TARGET = ROOT / "Ref/simulated_2b_correspondence/04_alternate_values_test"
ORIGINAL = ROOT / "Ref/simulated_2b_correspondence/03_missing_evidence_test/02_reference_do_not_upload"
VERSION = "20261008-alternate-values-1"
HOLDOUTS = ["photocopyRateAboveA3", "projectArchitectPhone", "siteInspectionEndDate", "quotationBeforeVariation"]
DISCLOSURE = "This is a fictional tender coordination exercise. All persons, organisations, proceedings and project decisions in this correspondence are invented. It has no contractual effect."


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def write_json(path: Path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def default_scenario():
    """Natural correspondence is authored independently of the original answers."""
    values = {k: None for k in read_json(ORIGINAL / "expected_values.json")["expectedInputs"]}
    evidence = {}
    records = []

    def record(n, date, title, kind="email"):
        item = {"recordId": f"SIM{n:02}", "date": date, "title": title, "kind": kind,
                "file": f"{n:02}_{date}_{kind}_{re.sub(r'[^a-z0-9]+', '_', title.lower()).strip('_')}.docx", "blocks": []}
        records.append(item)
        return item

    def para(rec, text, answers=None, heading=False):
        block = {"text": text, "heading": heading}
        rec["blocks"].append(block)
        for key, value in (answers or {}).items():
            if key in evidence:
                raise ValueError(f"Duplicate first disclosure: {key}")
            values[key] = value
            evidence[key] = {"recordId": rec["recordId"], "block": len(rec["blocks"]) - 1}

    r = record(1, "2026-04-06", "Request for tender coordination information")
    para(r, "Dear coordination team, please return the project decisions needed to prepare the NTT, SCT and SCC. We need the actual scope, pricing schedule, design responsibility, warranty coverage and inspection arrangements for this fictional project.")
    para(r, "This request confirms no project answer. If the communications do not settle a matter, leave it for the reviewer rather than using a standard-template default. We will assemble the replies in date order.")

    r = record(2, "2026-04-07", "Contract identity and combined foundation scope", "minutes")
    para(r, "The meeting confirmed Contract No. 20260247, Construction of Community Facilities at Tin Shui Wai Area 117. The scheme comprises a community hall, a district library, a sports building and an attached domestic accommodation block for resident caretakers.", {
        "contractTitle": {"number": "20260247", "title": "Construction of Community Facilities at Tin Shui Wai Area 117"}, "domesticBlocks": True, "projectInTinShuiWai": True})
    para(r, "This is a combined contract with foundation works. The appointed Contractor will design and construct the piles, pile caps and shallow footings for the buildings within the same contract, together with the building works.", {"foundationIncluded": True, "contractorDesignedFoundations": True})
    para(r, "All work locations and design details will be recorded in the subsequent replies. No inherited scope from another project is to be used.")

    r = record(3, "2026-04-08", "Programme duration and bond condition")
    para(r, "The accepted contract period is exactly 44 months. It therefore meets the condition for a contract period of 39 months or more. This programme duration is separate from the site inspection dates.", {"contractPeriodMonths": 44, "periodAtLeast39Months": True})
    para(r, "The foundation scope and this 44 month programme are the agreed facts for applying the combined foundation and duration condition in NTT10.")

    bills = [
        {"number": "1", "description": "Preliminaries", "type": "BQ"},
        {"number": "2", "description": "Preambles", "type": "BQ"},
        {"number": "3", "description": "Piled Foundations and Shallow Footings", "type": "BQ"},
        {"number": "4", "description": "Community Hall and Library Structure", "type": "BQ"},
        {"number": "5", "description": "Community Hall Finishes (All Provisional)", "type": "BQ"},
        {"number": "1", "description": "Ventilation Plant and Ductwork", "type": "SOR"},
        {"number": "2", "description": "Passenger Lifts and Escalators", "type": "SOR"},
    ]
    r = record(4, "2026-04-09", "Formal bill schedule for the community facilities")
    para(r, "Use the following complete formal Bill schedule of seven rows. BQ and SOR have separate number series, so BQ Bill 1 and SOR Bill 1 are distinct identities. The final words in parentheses in BQ Bill 5 are part of its formal description. The schedules have no eighth Bill.")
    bill_text = "Pricing type | Bill No. | Formal description\n" + "\n".join(f"{b['type']} | {b['number']} | {b['description']}" for b in bills)
    para(r, bill_text, {"billNos": bills, "pricingScheme": "BQ_SOR"})
    para(r, "The summary and collection sheets are support pages for these seven Bills. They do not create additional Bills or alter the formal descriptions above.")

    r = record(5, "2026-04-10", "Nominated subcontract arrangement confirmed", "minutes")
    para(r, "The adopted building services procurement arrangement is Nominated Sub-contracts (NSC). The ventilation plant and passenger lift packages will be nominated sub-contracts. Building Services Specialist Sub-contracts (BSSSC) are not adopted for this tender.", {"subcontractArrangement": "NSC"})
    para(r, "Retain the standard NTT13 nominated sub-contract provision. There is no BSSSC trade selection or replacement NTT13 wording to insert for this arrangement.")

    r = record(6, "2026-04-13", "Two envelope procedure and hardcopy issue")
    para(r, "Use the two-envelope tendering system. The technical submission and the priced offer are to be returned in separate sealed envelopes for this tender.", {"twoEnvelopeTendering": True})
    para(r, "The BQ issue format is Hardcopy. Tender documents are issued as printed originals; L10Pro electronic issue is not used. A reference PDF is a convenience copy and does not change the adopted hardcopy BQ issue format.", {"electronicTendering": "Hardcopy"})
    para(r, "No alternative tendering procedure is adopted. The photocopy rates for missing return copies will follow in the pricing reply.")

    r = record(7, "2026-04-14", "Foundation design and submission requirements")
    design = [
        {"component": "piling", "design": True, "execution": True, "scope": "Bored piles for the community hall and library"},
        {"component": "pilecaps", "design": True, "execution": True, "scope": "Pile caps for the community hall and library"},
        {"component": "footings", "design": True, "execution": True, "scope": "Shallow footings for the sports building and entrance canopy"},
        {"component": "other", "design": False, "execution": True, "scope": "Architect designed internal partitions"},
    ]
    para(r, "The complete component responsibility schedule is:\nComponent | Actual work scope | Contractor designs | Contractor executes\n" +
         "\n".join(f"{x['component']} | {x['scope']} | {'Yes' if x['design'] else 'No'} | {'Yes' if x['execution'] else 'No'}" for x in design), {"designResponsibilities": design})
    para(r, "The shallow footings serve the sports building and the main entrance canopy, which is a major external structure. Treat footings as serving buildings or major external structures.", {"footingsServeBuildingsOrMajorExternalStructures": True})
    para(r, "The foundation design will not be established at tender acceptance. The Contractor must submit its proposed foundation design with the Tender for assessment; subsequent checking and approval do not mean that the final design is already established at tender acceptance.", {"designEstablishedAtTenderAcceptance": False, "foundationDesignSubmittedWithTender": True})
    para(r, "RSE endorsed submissions are required. The additional submission requirements are:\nSubmit a construction noise management proposal with the Tender.\nSubmit a traffic diversion method statement with the Tender.", {"rseSubmissionsRequired": True, "additionalSubmissions": [{"text": "Submit a construction noise management proposal with the Tender."}, {"text": "Submit a traffic diversion method statement with the Tender."}]})
    para(r, "Should the special requirement to obtain a design-related quotation before instructing a Variation be adopted? This reply contains no adopted decision on that requirement. Keep the matter open for the business reviewer.")

    r = record(8, "2026-04-15", "Railway protection and valuable tree requirements", "minutes")
    para(r, "Works extend into a railway protection area. Railway protection plans are not available for inspection at this stage, so do not state that those plans are available.", {"railwayProtectionAreaWorks": True, "railwayProtectionPlansAvailable": False})
    trees = [{"serial": "OVT-TSW-117-03"}, {"serial": "OVT-TSW-117-08"}]
    para(r, "The complete Old or Valuable Trees register for the site contains two distinct serials: OVT-TSW-117-03 and OVT-TSW-117-08. Both trees are to remain protected in place.", {"oldValuableTrees": trees})
    tree_wording = "Provide fenced protection zones around OVT-TSW-117-03 and OVT-TSW-117-08. No storage, vehicle movement or excavation is permitted within those zones without the Architect's written approval."
    para(r, "Adopt the following project wording for the two Old or Valuable Trees:\n" + tree_wording, {"ovtProjectSpecificText": tree_wording})

    r = record(9, "2026-04-16", "Works scope and construction method reply")
    para(r, "Works Subject to Excision are specified: the optional covered link between the library and sports building may be excised. Site formation works are included, and works in the public road require an Excavation Permit.", {"worksSubjectToExcision": True, "siteFormationWorks": True, "excavationPermitRequired": True})
    para(r, "No Subsidised Sale Flats production is included. Neither exhibition flats nor show flats are required; the residential element is caretaker accommodation only.", {"subsidisedSaleFlatsProduction": False, "showFlatsRequirement": "none"})
    para(r, "Construction of the domestic caretaker accommodation block is included. CON8 volumetric precast components are not used for that block. No PSE submission for CON8 volumetric precast components is therefore requested.", {"volumetricPrecastComponents": False})
    para(r, "The Tender is not based on precast concrete facades. Precast facades are neither required nor permitted to be proposed under APE2.A410 for this project.", {"precastFacadeTenderBasis": False, "precastFacadePermission": False})
    para(r, "The WTO Government Procurement Agreement does not apply to this fictional tender.", {"wtoGpaApplies": False})

    r = record(10, "2026-04-17", "Pricing basis and photocopy charge reply")
    para(r, "The adopted pricing scheme combines Bills of Quantities and Schedules of Rates (BQ plus SOR). All quantities in the BQ Bills are provisional. The two SOR Bills remain rate schedules, not measured BQ Bills.", {"allBqQuantitiesProvisional": True})
    para(r, "The seven formal Bill identities are the schedule already supplied. BQ and SOR numbering are separate. Do not collapse identically numbered Bills across those pricing types.")
    para(r, "Where the required number of tender copies is missing, the photocopy charge for each document page of or below A3 size is HK$12.8.", {"photocopyRateUpToA3": 12.8})
    para(r, "What photocopy charge per page applies above A3? No project rate for pages exceeding A3 has been authorised in the supplied communications. Leave the rate for the reviewer to determine.")
    para(r, "Both SOR detailed schedules and subsequent addenda are issued in hardcopy. The reference PDF does not create an additional Bill identity.")

    r = record(11, "2026-04-20", "Installation and waterproofing warranty scope", "minutes")
    para(r, "The project and its domestic caretaker block are in Tin Shui Wai. Roofing installation is included, but a separate roofing warranty is not required for this contract.", {"roofingInstallationIncluded": True, "roofingWarrantyRequired": False})
    para(r, "Refuge floor installation is included in the sports building and its installation warranty is required.", {"refugeFloorInstallationIncluded": True, "refugeFloorWarrantyRequired": True})
    para(r, "A warranty for the other waterproofing installations is required. The following specification list is complete and has nine items. Each continuation line belongs to the preceding numbered item.", {"otherWaterproofingWarrantyRequired": True})
    water = "\n".join([
        "1. Changing room shower floors and kerbs, including the recessed accessible shower bays, waterproofed bench supports and pipe penetrations. Turn the membrane up behind the wall lining, seal the interfaces with the floor drains and continue the protection beneath the entrance thresholds;",
        "2. Library archive basement retaining walls, including wall to slab junctions, construction joints and service sleeves. The tanking must extend around the recessed book delivery entrance and terminate at the approved external ground level, with compatible seals where the wall meets the loading apron;",
        "3. Sports hall equipment wash bays and their trapped drains, including cleaning cupboard floors and the wash down channels. Provide continuous waterproofing beneath equipment plinths and around the floor gully collars, and protect the membrane where maintenance trolleys pass between the wet bay and storage room;",
        "4. Community hall catering preparation sinks, service trenches and splashback upstands, including the dishwashing recess and the small refrigerated store. Continue waterproofing through the trench junctions and under the counter legs, with sealed access covers and membrane returns at every wet area door;",
        "5. Rainwater harvesting tank external faces, inspection sumps and connecting channels, including the tank base joints, pipe connection chambers and overflow route to the landscape drain;\n   Include the waterproofed junctions at maintenance access openings and ensure that the protection boards do not obstruct the inspection or replacement of the access seals;",
        "6. Courtyard fountain basin, overflow chamber and adjacent service duct, including the fountain pump recess and the cable entry sleeves. Waterproof the full basin wall and floor surfaces beneath the finishes, return the system over the perimeter edge and seal the movement joints separating the basin from the paved courtyard;",
        "7. Loading bay ramps over the underground library store and associated movement joints, including the wheel guide plinths and drainage channel at the ramp foot. The waterproofing must continue behind the kerbs and around the store ventilation openings, with a traffic compatible wearing layer over all vehicle routes;",
        "8. Caretaker block laundry balconies and clothes drying platforms, including the recessed washing machine positions and the shared laundry entrance;\n   Include membrane returns around thresholds and drainage outlets, seal balustrade post fixings and continue the waterproofed surface beneath the removable service platforms used to maintain the washing machine supply pipes;",
        "9. Green roof planting troughs over the community hall foyer and their perimeter upstands, including the irrigation valve pockets and emergency overflow outlets. The waterproofing must include compatible root resistant protection, continuous seals at trough corners and accessible inspection zones around every drain outlet."
    ])
    para(r, water, {"otherWaterproofingSpecificationAreas": water})
    para(r, "No acrylic paint warranty and no wall tiling warranty are required. Play equipment and Impact Absorbing Surfacing Materials are not included. An epoxy coated cast iron pipe installation warranty is not required.", {"acrylicPaintWarrantyRequired": False, "wallTilingWarrantyRequired": False, "playEquipmentIncluded": False, "iasmIncluded": False, "epoxyCastIronPipeWarrantyRequired": False})

    r = record(12, "2026-04-21", "Project architect and inspection arrangements reply")
    para(r, "For SCT7, the Project Architect post is Senior Architect Community Projects. Use the professional title Dr rather than Mr or Ms. The appointed fictional officer's name is Julian K. M. LEE.", {"projectArchitectPost": "Senior Architect Community Projects", "projectArchitectSalutation": "Other", "projectArchitectOtherTitle": "Dr", "projectArchitectName": "Julian K. M. LEE"})
    para(r, "What telephone number is to be published for the Project Architect? No telephone number has been confirmed. Contact the fictional coordination mailbox architect@consense-scenario.invalid meanwhile; this mailbox is not a telephone number.")
    para(r, "The Specification Library is available for inspection at floor 3, Block C, Fictional Civic Project Office, 88 Learning Road, Tin Shui Wai.", {"specificationInspectionFloor": "3", "specificationInspectionBlock": "C"})
    para(r, "The drawings are available for inspection at floor 5, Block D, Fictional Civic Project Office, 88 Learning Road, Tin Shui Wai.", {"drawingsInspectionFloor": "5", "drawingsInspectionBlock": "D"})
    para(r, "The site inspection start date is 2026-06-08. The inspection end date has not been set and must remain unanswered; the start date does not authorise a guessed end date.", {"siteInspectionStartDate": "2026-06-08"})
    restrictions = [{"text": "Apply in writing at least ten working days before the proposed visit."}, {"text": "Each visiting group is limited to six persons and must wear safety helmets and protective footwear."}]
    para(r, "The actual site visit restrictions are:\n" + "\n".join(x["text"] for x in restrictions), {"siteVisitRestrictions": restrictions})

    r = record(13, "2026-04-22", "Sections and special clause decisions", "minutes")
    sections = [{"designation": "Section A", "workTypes": ["building", "foundation"], "location": "Community hall and library at the north site"}, {"designation": "Section B", "workTypes": ["demolition"], "location": "Former workshop at the south site"}, {"designation": "Section C", "workTypes": ["building", "foundation"], "location": "Sports building and caretaker block at the east site"}]
    para(r, "Adopt these three Sections:\nSection designation | Works types | Site or location\n" + "\n".join(f"{s['designation']} | {', '.join(s['workTypes'])} | {s['location']}" for s in sections), {"sections": sections})
    para(r, "The building sites and demolition site are separated: building work is at the north and east sites, while demolition is at the south site.", {"buildingDemolitionSitesSeparated": True})
    para(r, "Adopt SCC8.303 for weather extension of time. The Section comprising demolition works only is Section B.", {"weatherEotClause": "SCC8.303", "demolitionOnlySection": "Section B"})
    para(r, "Adopt the late possession provision. Also adopt the additional loss provision: the specified Sections are Section A and Section C, and the relevant Section is Section C.", {"latePossessionAdopted": True, "additionalLossAdopted": True, "additionalLossSpecifiedSections": "Section A and Section C", "additionalLossRelevantSection": "Section C"})

    r = record(14, "2026-04-23", "Police training and contractor legal form decisions", "minutes")
    para(r, "Adopt the police training clause for the demolition site. The former workshop is a standalone building, and the south site is suitable for police training during the first month of the contract.", {"policeTrainingAdopted": True, "demolitionStandaloneBuilding": True, "siteSuitableForTrainingFirstMonth": True})
    para(r, "The Contractor legal form is Other: an ordinary incorporated company, not a partnership and not an incorporated or unincorporated joint venture. The legal form is known; the fallback choices for an unknown entity form are not needed.", {"contractorLegalForm": "other"})
    para(r, "The design related quotation before Variation decision is still outstanding. No approval or rejection of that optional requirement was recorded here.")

    r = record(15, "2026-04-24", "Advance payment and final information review", "minutes")
    para(r, "Adopt the advance payment arrangement under SCC14.305 for this draft. Set a project specific repayment period of eight consecutive months, with the first repayment instalment deducted in the fifth interim payment certificate. These are express overrides for this fictional project.", {"advancePaymentAdopted": True, "advanceRepaymentMonths": 8, "advanceRepaymentFirstCertificate": 5})
    para(r, "Four decisions remain unconfirmed: the photocopy rate for pages above A3, the Project Architect telephone number, the site inspection end date, and the optional design related quotation before Variation requirement. The communications contain no adopted values for them; keep these questions available for the reviewer.")
    para(r, "The replies for all other adopted matters remain effective. This review introduces no additional Bills, no extra waterproofing areas and no revised Section designation.")
    return {"fixtureVersion": VERSION, "synthetic": True, "records": records, "values": values, "evidence": evidence,
            "withheld": HOLDOUTS, "prohibitedLegacyPhrases": ["20250101", "Tung Chung Area 98", "Emily T.L. WONG", "2761 6161", "2025-02-17", "2025-02-21", "Bill No. 12.2", "Bill Nos. 9 to 11"]}


def applicability(field, values):
    terms = field.get("condition", {}).get("all", [])
    results = []
    for term in terms:
        value = values[term["field"]]
        op, expected = term["operator"], term.get("value")
        if op == "unknown":
            results.append(value is None)
        elif value is None:
            results.append(None)
        elif op == "eq":
            results.append(value == expected)
        elif op == "in":
            results.append(value in expected)
        elif op == "countGt":
            results.append(len(value) > expected)
        elif op == "includesComponent":
            results.append(any(row.get("component") == expected for row in value))
        else:
            raise ValueError(f"Unsupported catalogue condition {op}")
    if False in results:
        return "inactive"
    return "prerequisite_unknown" if None in results else "active"


def canonical_zip(path):
    """Stable OOXML bytes for an unchanged scenario, including ZIP timestamps."""
    with ZipFile(path) as archive:
        members = [(name, archive.read(name)) for name in archive.namelist()]
    from zipfile import ZipInfo
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in members:
            info = ZipInfo(name, date_time=(2026, 4, 6, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 0
            archive.writestr(info, content)


def author_doc(record, path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.75)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if name == "Normal" else 17 if name == "Title" else 12)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.08
        ppr = style.element.find(qn("w:pPr"))
        if ppr is not None:
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)
    doc.styles["Title"].paragraph_format.space_after = Pt(12)
    doc.core_properties.title = record["title"]
    doc.core_properties.author = "ConSense fictional coordination team"
    doc.core_properties.created = datetime(2026, 4, 6)
    doc.core_properties.modified = datetime(2026, 4, 6)
    doc.add_paragraph(record["title"], "Title")
    doc.add_paragraph(DISCLOSURE)
    doc.add_paragraph(f"Record {record['recordId']}    Date {record['date']}    {'Meeting minutes' if record['kind'] == 'minutes' else 'Email reply'}")
    if record["kind"] == "minutes":
        doc.add_paragraph("Attendees: Fictional Project Architect, Structural Engineer, Procurement Officer and Quantity Surveyor. Recorder: Fictional Quantity Surveyor.")
    else:
        doc.add_paragraph("From: Fictional Quantity Surveyor <qs@consense-scenario.invalid>\nTo: Fictional Coordination Team <team@consense-scenario.invalid>")
    block_to_paragraph = {}
    for index, block in enumerate(record["blocks"]):
        block_to_paragraph[index] = len(doc.paragraphs)
        p = doc.add_paragraph(block["text"], "Heading 1" if block.get("heading") else "Normal")
        p.paragraph_format.widow_control = True
    if record["kind"] == "email":
        doc.add_paragraph("Regards,\nFictional Quantity Surveyor")
    doc.save(path)
    canonical_zip(path)
    return block_to_paragraph


def check_value(field, value, prefix):
    errors = []
    if value is None:
        return errors
    kind = field["kind"]
    options = [o["value"] for o in field.get("options", [])]
    if kind == "boolean" and type(value) is not bool:
        errors.append(f"{prefix}: expected boolean")
    if kind == "number" and (type(value) not in (int, float)):
        errors.append(f"{prefix}: expected number")
    if kind == "choice" and value not in options:
        errors.append(f"{prefix}: illegal choice {value}")
    if kind == "multiselect" and (not isinstance(value, list) or any(v not in options for v in value)):
        errors.append(f"{prefix}: illegal multiselect")
    if kind in ("date", "text") and not isinstance(value, str):
        errors.append(f"{prefix}: expected string")
    if kind == "date" and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        errors.append(f"{prefix}: illegal ISO date")
    if kind in ("list", "contract"):
        rows = value if kind == "list" else [value]
        if not isinstance(rows, list) or any(not isinstance(r, dict) for r in rows):
            return errors + [f"{prefix}: expected objects"]
        for index, row in enumerate(rows):
            columns = {c["key"]: c for c in field.get("columnFields", [])}
            errors.extend(f"{prefix}[{index}]: unknown column {c}" for c in set(row) - set(columns))
            for key, column in columns.items():
                if key not in row and not column.get("optional", False):
                    errors.append(f"{prefix}[{index}]: missing required column {key}")
                errors.extend(check_value(column, row.get(key), f"{prefix}[{index}].{key}"))
    return errors


def build(target, scenario, reuse_documents=False):
    upload, ref = target / "01_upload_documents", target / "02_reference_do_not_upload"
    upload.mkdir(parents=True, exist_ok=True)
    ref.mkdir(parents=True, exist_ok=True)
    catalogue = read_json(ORIGINAL / "input_catalogue_snapshot.json")
    fields = {f["key"]: f for g in catalogue["data"]["groups"] for f in g["fields"]}
    paragraph_maps, record_map = {}, {}
    for record in scenario["records"]:
        path = upload / record["file"]
        if reuse_documents:
            existing = Document(path)
            texts = [p.text for p in existing.paragraphs]
            matches = {}
            for index, block in enumerate(record["blocks"]):
                slots = [i for i, text in enumerate(texts) if text == block["text"]]
                if len(slots) != 1:
                    raise ValueError(f"{record['recordId']}: cannot uniquely rebind source paragraph")
                matches[index] = slots[0]
            paragraph_maps[record["recordId"]] = matches
        else:
            paragraph_maps[record["recordId"]] = author_doc(record, path)
        doc = Document(path)
        (ref / Path(record["file"]).with_suffix(".txt")).write_text("\n\n".join(p.text for p in doc.paragraphs) + "\n", encoding="utf-8")
        record_map[record["recordId"]] = record
    expected = {}
    for key, value in scenario["values"].items():
        active = applicability(fields[key], scenario["values"])
        entry = {"value": value, "status": "not_applicable" if active == "inactive" else "unknown" if key in HOLDOUTS else "derived" if key == "periodAtLeast39Months" else "provided",
                 "origin": "deliberately_withheld" if key in HOLDOUTS else "inactive_detail" if active == "inactive" else "authored_test_decision",
                 "rationale": "The applicable project value is deliberately absent; no default or negative answer is authorised." if key in HOLDOUTS else "This detail is inactive under the final adopted prerequisite." if active == "inactive" else "The following exact paragraph is an explicit fictional project decision; it is not a real 2b source fact.",
                 "firstDisclosure": None, "sourceReferences": [], "humanAdoptedInApplication": False, "applicability": active}
        if key in scenario["evidence"]:
            slot = scenario["evidence"][key]
            record = record_map[slot["recordId"]]
            entry["firstDisclosure"] = slot["recordId"]
            entry["inputEvidence"] = {"recordId": slot["recordId"], "file": record["file"], "nativeParagraph": paragraph_maps[slot["recordId"]][slot["block"]] + 1, "quote": record["blocks"][slot["block"]]["text"], "sha256": sha(upload / record["file"])}
        expected[key] = entry
    oracle = {"schemaVersion": 2, "fixtureVersion": VERSION, "synthetic": True,
              "case": "Wholly fictional alternate community facilities tender", "purpose": "Held out values and changed branches test the frozen Drafting harness without uploading the answer key.",
              "catalogueRuleVersion": catalogue["data"]["ruleVersion"], "groupCount": len(catalogue["data"]["groups"]), "inputCount": len(fields), "expectedInputs": expected,
              "derivedDecisions": {"g1aJointCondition": {"value": True, "dependencies": ["foundationIncluded", "periodAtLeast39Months"], "rule": "foundationIncluded AND periodAtLeast39Months"}, "nscApplicable": {"value": True, "dependencies": ["subcontractArrangement"]}, "oldValuableTreeCount": {"value": 2, "dependencies": ["oldValuableTrees"]}},
              "sourceContext": {"acceptedContractPeriodMonths": 44, "actualReturnMedia": "two_sealed_hardcopy_envelopes"},
              "effectiveTerms": {}, "billDistributionAndUsageOracle": {"caseSpecific": True, "sorDetailedIssue": {"numbers": ["1", "2"], "originalTender": "hardcopy", "addenda": "hardcopy", "rationale": "Explicit alternate issue decision in SIM10"}},
              "knownRepresentationLimitations": [{"topic": "Pricing namespaces", "caseTruth": "BQ Bills 1-5 and SOR Bills 1-2 are seven distinct Bills", "limitation": "Bill number alone is not a unique identity; pricing type must be preserved"}],
              "excluded": ["Tender A/B", "bid prices", "future performance", "hidden targetOverrides without explicit human action"],
              "scoring": {"modelOutputIsUnconfirmedPrefill": True, "falseAndEmptyAreAnswers": True, "inactiveDetailsMustNotBeInvented": True, "compareBusinessMeaningNotRandomInternalIds": True, "caseSourceAndAuthoredTestDecisionMustStaySeparate": True, "absentRequiredFactsAreNotNo": True, "deliberatelyWithheldInputsMustRemainNull": HOLDOUTS, "unknownMustNotBecomeFalseZeroEmptyStringOrDefaultDate": True},
              "variant": "alternate_values_with_four_active_holdouts",
              "gradingMetadata": {"examTitle": "不同变量值的模拟通信稳定性考试", "candidateExamScope": "79 inputs, 15 fictional communications, seven type-qualified Bills and nine complete waterproofing areas", "additionalBusinessBoundary": "Additional distribution and numbering observations are separate from input candidate correctness and human adoption.", "additionalBusinessNotes": ["BQ1 and SOR1, and BQ2 and SOR2, are distinct parent Bills.", "The formal BQ5 description retains its All Provisional qualifier; this does not create a new Bill.", "The waterproofing specification has exactly nine items, with continuation lines belonging to items 5 and 8.", "The reference answers were not uploaded; all output remains unconfirmed suggestions."]}}
    ledger, cumulative = [], {k: None for k in fields}
    for record in scenario["records"]:
        answers = {k: scenario["values"][k] for k, slot in scenario["evidence"].items() if slot["recordId"] == record["recordId"]}
        cumulative.update(answers)
        ledger.append({"recordId": record["recordId"], "file": record["file"], "newAnswers": answers, "newAnswerOrigins": {k: "authored_test_decision" for k in answers}, "cumulativeValues": copy.deepcopy(cumulative), "applicability": {k: applicability(field, cumulative) for k, field in fields.items()}, "providedCount": sum(v is not None for v in cumulative.values()), "automaticHumanAdoptionExpected": False})
    write_json(ref / "expected_values.json", oracle)
    write_json(ref / "incremental_expected_states.json", ledger)
    write_json(ref / "input_catalogue_snapshot.json", catalogue)
    write_json(ref / "scenario_config.json", scenario)
    write_json(ref / "withheld_control.json", {"variant": "four_active_questions_without_answers", "doNotUpload": True, "holdouts": {key: {"expectedBehavior": "null/unknown; available for human determination; never infer false, zero or default"} for key in HOLDOUTS}})
    old = read_json(ORIGINAL / "expected_values.json")["expectedInputs"]
    changed = [{"key": k, "oldValue": old[k]["value"], "newValue": v, "oldStatus": old[k]["status"], "newStatus": expected[k]["status"], "oldApplicability": old[k]["applicability"], "newApplicability": expected[k]["applicability"]} for k, v in scenario["values"].items() if v != old[k]["value"] or old[k]["status"] != expected[k]["status"] or old[k]["applicability"] != expected[k]["applicability"]]
    write_json(ref / "change_ledger.json", {"baselineFixture": "03_missing_evidence_test", "changedBusinessValues": sum(scenario["values"][k] != old[k]["value"] for k in fields), "changedEntries": changed})
    lines = ["# Alternate fictional correspondence fixture", "", "Upload only the 15 DOCX files in `01_upload_documents`. Every file is fictional and carries its own disclosure. The reference directory contains the answer key, exact source paragraphs, scenario configuration and cumulative stage ledger; do not upload that directory.", "", "This case changes the project identity, foundation scope, programme, subcontract model, format, design responsibilities, railway and tree branches, pricing lists, warranties, architect, locations, inspection arrangements, Sections and special clauses. Bills use two independent pricing namespaces: BQ1–5 and SOR1–2. Waterproofing has nine complete items, including multiline continuations.", "", "Four active facts are deliberately absent: the above-A3 photocopy rate, architect telephone, inspection end date and design quotation before Variation decision. Absence must remain unanswered. Existing candidate results are suggestions, never automatic human adoption.", "", "Reproduction uses the bundled Python runtime with python-docx:", "", "```powershell", "& 'C:/Users/Administrator/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/python.exe' scripts/generate-alternate-drafting-fixture.py --generate", "& 'C:/Users/Administrator/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/python.exe' scripts/generate-alternate-drafting-fixture.py --validate", "```", "", "Regeneration is deterministic and deliberately refuses a nonempty upload directory unless `--overwrite-generated` is specified. The original 03 fixture, production code, prompts and runtime are untouched. Rendered QA remains outside the upload directory.", ""]
    (target / "README.md").write_text("\n".join(lines), encoding="utf-8")


def validate(target):
    upload, ref = target / "01_upload_documents", target / "02_reference_do_not_upload"
    scenario, oracle = read_json(ref / "scenario_config.json"), read_json(ref / "expected_values.json")
    ledger, catalogue = read_json(ref / "incremental_expected_states.json"), read_json(ref / "input_catalogue_snapshot.json")
    fields = {f["key"]: f for g in catalogue["data"]["groups"] for f in g["fields"]}
    errors = []
    uploads = list(upload.iterdir())
    if len(uploads) != 15 or any(p.suffix != ".docx" for p in uploads):
        errors.append("Upload directory must contain exactly 15 DOCX files and no oracle or QA files")
    if len(fields) != 79 or set(oracle["expectedInputs"]) != set(fields):
        errors.append("Expected oracle must exactly cover all 79 catalogue inputs")
    combined = "\n".join(p.text for file in uploads if file.suffix == ".docx" for p in Document(file).paragraphs)
    errors.extend(f"Legacy phrase leaked into new source: {phrase}" for phrase in scenario["prohibitedLegacyPhrases"] if phrase in combined)
    errors.extend(f"Machine variable key leaked into uploaded body: {key}" for key in fields if re.search(r"\b" + re.escape(key) + r"\b", combined))
    mailboxes = re.findall(r"[\w.+-]+@[\w.-]+", combined)
    if any(not mailbox.endswith(".invalid") for mailbox in mailboxes):
        errors.append("All fictional contacts must use the .invalid domain")
    for record in scenario["records"]:
        doc = Document(upload / record["file"])
        if doc.tables:
            errors.append(f"{record['recordId']}: expected native paragraph data, not nested Word tables")
        if doc.paragraphs[0].style.name != "Title" or DISCLOSURE not in [p.text for p in doc.paragraphs]:
            errors.append(f"{record['recordId']}: missing Title style or fictional disclosure")
        if doc.styles["Title"].element.findall(".//" + qn("w:pBdr")):
            errors.append(f"{record['recordId']}: title must not carry an inherited paragraph border")
    for key, entry in oracle["expectedInputs"].items():
        errors.extend(check_value(fields[key], entry["value"], key))
        actual_app = applicability(fields[key], scenario["values"])
        if actual_app != entry["applicability"]:
            errors.append(f"{key}: applicability does not match catalogue")
        if actual_app == "inactive" and (entry["value"] is not None or entry["status"] != "not_applicable"):
            errors.append(f"{key}: inactive value must be absent")
        if entry["value"] is None and actual_app == "active" and key not in HOLDOUTS:
            errors.append(f"{key}: unexpected active missing answer outside the four holdouts")
        if entry.get("sourceReferences") != []:
            errors.append(f"{key}: fictional value cannot cite real source documents")
        if entry["value"] is not None:
            ev = entry.get("inputEvidence", {})
            file = upload / ev.get("file", "missing")
            if not file.is_file() or ev.get("sha256") != sha(file):
                errors.append(f"{key}: source SHA missing or inconsistent")
                continue
            paragraphs = Document(file).paragraphs
            ix = ev.get("nativeParagraph", 0) - 1
            if ix < 0 or ix >= len(paragraphs) or paragraphs[ix].text != ev.get("quote"):
                errors.append(f"{key}: quote must equal a complete native paragraph")
            if entry.get("firstDisclosure") != ev.get("recordId"):
                errors.append(f"{key}: first disclosure and source record mismatch")
        elif key in HOLDOUTS and (entry["status"] != "unknown" or actual_app != "active" or entry.get("firstDisclosure") is not None):
            errors.append(f"{key}: deliberately withheld fact must stay active and unknown")
    cumulative = {k: None for k in fields}
    if len(ledger) != 15:
        errors.append("Expected ledger must contain 15 stages")
    for record, stage in zip(scenario["records"], ledger):
        intended = {key: e["value"] for key, e in oracle["expectedInputs"].items() if e.get("firstDisclosure") == record["recordId"]}
        if stage["recordId"] != record["recordId"] or stage["newAnswers"] != intended:
            errors.append(f"{record['recordId']}: incorrect new answer disclosure")
        cumulative.update(intended)
        if stage["cumulativeValues"] != cumulative:
            errors.append(f"{record['recordId']}: future-answer leak or incorrect cumulative values")
        expected_app = {k: applicability(field, cumulative) for k, field in fields.items()}
        if stage["applicability"] != expected_app:
            errors.append(f"{record['recordId']}: incorrect incremental applicability")
        if stage["providedCount"] != sum(v is not None for v in cumulative.values()):
            errors.append(f"{record['recordId']}: incorrect provided count")
        if any(cumulative[k] is not None for k in HOLDOUTS):
            errors.append(f"{record['recordId']}: withheld answer leak")
    if cumulative != scenario["values"]:
        errors.append("Final stage does not equal all final expected values")
    bills = scenario["values"]["billNos"]
    bill_ids = [(b["type"], b["number"]) for b in bills]
    if len(bills) != 7 or len(set(bill_ids)) != 7 or set(bill_ids) != {("BQ", str(n)) for n in range(1, 6)} | {("SOR", "1"), ("SOR", "2")}:
        errors.append("Expected seven explicit BQ/SOR namespace identities")
    water = scenario["values"]["otherWaterproofingSpecificationAreas"]
    if re.findall(r"(?m)^(\d+)\.", water) != [str(n) for n in range(1, 10)]:
        errors.append("Waterproofing must contain all numbered items 1-9, with continuations preserved")
    old = read_json(ORIGINAL / "expected_values.json")["expectedInputs"]
    changed = sum(scenario["values"][k] != old[k]["value"] for k in fields)
    if changed < 35:
        errors.append("Alternate scenario must change at least 35 business values")
    if scenario["values"]["periodAtLeast39Months"] != (scenario["values"]["contractPeriodMonths"] >= 39):
        errors.append("44 month programme threshold is inconsistent")
    section_by_name = {s["designation"]: s for s in scenario["values"]["sections"]}
    demolition_name = scenario["values"]["demolitionOnlySection"]
    if demolition_name not in section_by_name or section_by_name[demolition_name]["workTypes"] != ["demolition"]:
        errors.append("Demolition-only designation must identify the actual demolition-only Section")
    if scenario["values"]["additionalLossRelevantSection"] not in section_by_name:
        errors.append("Additional loss relevant Section must exist in the adopted schedule")
    serials = [t["serial"] for t in scenario["values"]["oldValuableTrees"]]
    if len(serials) != len(set(serials)) or len(serials) != 2:
        errors.append("The two OVT records must have distinct nonblank serials")
    report = {"verdict": "FAIL" if errors else "PASS", "fixtureVersion": VERSION,
              "documentCount": len(uploads), "inputCount": len(fields), "changedBusinessValues": changed,
              "statusCounts": dict(Counter(e["status"] for e in oracle["expectedInputs"].values())),
              "applicabilityCounts": dict(Counter(e["applicability"] for e in oracle["expectedInputs"].values())),
              "providedOrDerivedCount": sum(e["value"] is not None for e in oracle["expectedInputs"].values()),
              "expectedNullCount": sum(e["value"] is None for e in oracle["expectedInputs"].values()),
              "activeWithheldCount": len(HOLDOUTS), "billCount": len(bills), "waterproofingItemCount": 9,
              "errors": errors, "sourceFiles": [{"file": p.name, "sha256": sha(p)} for p in sorted(uploads)],
              "oracleSha256": sha(ref / "expected_values.json"), "scenarioSha256": sha(ref / "scenario_config.json"), "ledgerSha256": sha(ref / "incremental_expected_states.json")}
    write_json(ref / "validation.json", report)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--target", type=Path, default=DEFAULT_TARGET)
    parser.add_argument("--generate", action="store_true")
    parser.add_argument("--validate", action="store_true")
    parser.add_argument("--rebind-metadata", action="store_true", help="Rebuild sealed oracle anchors against byte-identical existing DOCX sources")
    parser.add_argument("--overwrite-generated", action="store_true")
    parser.add_argument("--scenario", type=Path)
    args = parser.parse_args()
    target = args.target.resolve()
    if args.generate:
        uploads = target / "01_upload_documents"
        if uploads.exists() and any(uploads.iterdir()) and not args.overwrite_generated:
            parser.error("Refuse to replace existing DOCX sources; use --overwrite-generated explicitly")
        scenario = read_json(args.scenario) if args.scenario else default_scenario()
        if len(scenario["records"]) != 15:
            parser.error("Exactly 15 authored correspondence records are required")
        build(target, scenario)
    if args.rebind_metadata:
        scenario = read_json(args.scenario) if args.scenario else default_scenario()
        build(target, scenario, reuse_documents=True)
    if args.generate or args.validate or args.rebind_metadata:
        report = validate(target)
        print(json.dumps({k: v for k, v in report.items() if k != "sourceFiles"}, ensure_ascii=False, indent=2))
        raise SystemExit(0 if report["verdict"] == "PASS" else 1)
    parser.error("Choose --generate or --validate; this script never starts inference")


if __name__ == "__main__":
    main()
